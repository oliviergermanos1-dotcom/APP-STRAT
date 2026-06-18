#!/usr/bin/env python3
"""Parse the DSM Import 2025 Excel into a structured JSON consumable by the PPTX generator.

Input  : examples/dsm_import_2025.xlsx
Output : lib/pptx/data/dsm.json
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "examples" / "dsm_import_2025.xlsx"
OUT  = ROOT / "lib" / "pptx" / "data" / "dsm.json"

MONTHS_FR = ["Janvier","Février","Mars","Avril","Mai","Juin",
             "Juillet","Août","Septembre","Octobre","Novembre","Décembre"]


def find_header_row(ws, label):
    """Return the row index whose first cell starts with `label`."""
    for r in range(1, min(ws.max_row, 20) + 1):
        v = ws.cell(row=r, column=1).value
        if v and isinstance(v, str) and v.strip().lower().startswith(label.lower()):
            return r
    return None


def to_float(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


AGGREGATE_LABELS = {"total général", "total general", "grand total", "total"}


def is_aggregate(label):
    return label.strip().lower() in AGGREGATE_LABELS


def parse_monthly_sheet(ws, label_col_name):
    """Parse a sheet with structure: [label] | 12 mois | Total général."""
    hdr = find_header_row(ws, label_col_name)
    if hdr is None:
        raise ValueError(f"Header '{label_col_name}' not found in sheet '{ws.title}'")
    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if not label or not isinstance(label, str):
            continue
        if is_aggregate(label):
            continue
        monthly = [to_float(ws.cell(row=r, column=c).value) for c in range(2, 14)]
        total = to_float(ws.cell(row=r, column=14).value)
        if total == 0 and sum(monthly) == 0:
            continue
        rows.append({
            "name": label.strip(),
            "monthly": monthly,
            "total": total or sum(monthly),
        })
    rows.sort(key=lambda x: x["total"], reverse=True)
    return rows


def parse_pol(ws):
    hdr = find_header_row(ws, "Port de chargement")
    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if not label or not isinstance(label, str):
            continue
        if is_aggregate(label):
            continue
        monthly = [to_float(ws.cell(row=r, column=c).value) for c in range(2, 14)]
        total = to_float(ws.cell(row=r, column=14).value)
        if total == 0 and sum(monthly) == 0:
            continue
        rows.append({"port": label.strip(), "monthly": monthly, "total": total or sum(monthly)})
    rows.sort(key=lambda x: x["total"], reverse=True)
    return rows


def compute_pdm(rows, top_n=10):
    """Return top N rows with computed pdm (% of total market)."""
    market = sum(r["total"] for r in rows)
    top = rows[:top_n]
    for r in top:
        r["pdm"] = round((r["total"] / market) * 100, 2) if market else 0
    return {"market_total": round(market, 2), "top": top}


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    armateurs = parse_monthly_sheet(wb["armateur au bl"], "Armateur BL")
    manutent  = parse_monthly_sheet(wb["manutentionnaire"], "Manutentionaire")
    consign   = parse_monthly_sheet(wb["consignaitaire"], "Consignataire")
    pol       = parse_pol(wb["POL"])

    result = {
        "source": "BESOIN_RAPPORT_DSM_ABJ_ET_SPY.xlsx",
        "period": "Import 2025 hors PP — Tonnage tous conditionnements",
        "armateurs":  compute_pdm(armateurs, top_n=10),
        "manutentionnaires": compute_pdm(manutent, top_n=10),
        "consignataires":    compute_pdm(consign,  top_n=10),
        "pol_top20":         compute_pdm(pol,      top_n=20),
        "months_fr": MONTHS_FR,
    }

    # AGL position (computed)
    for key in ["armateurs", "manutentionnaires", "consignataires"]:
        block = result[key]
        agl_row = next(
            (i for i, r in enumerate(armateurs if key == "armateurs"
                                     else manutent if key == "manutentionnaires"
                                     else consign)
             if "AGL" in r["name"].upper()),
            None,
        )
        if agl_row is not None:
            full = (armateurs if key == "armateurs"
                    else manutent if key == "manutentionnaires"
                    else consign)
            block["agl_rank"] = agl_row + 1
            block["agl_total"] = full[agl_row]["total"]
            block["agl_pdm"] = round((full[agl_row]["total"] / block["market_total"]) * 100, 2)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(f"Wrote {OUT}")
    print(f"  Armateurs market : {result['armateurs']['market_total']:,.0f} T")
    print(f"  Manutent. market : {result['manutentionnaires']['market_total']:,.0f} T")
    print(f"  Consign. market  : {result['consignataires']['market_total']:,.0f} T")
    print(f"  AGL armateur rank: {result['armateurs'].get('agl_rank', 'N/A')}")
    print(f"  AGL manut. rank  : {result['manutentionnaires'].get('agl_rank', 'N/A')}")
    print(f"  AGL consign rank : {result['consignataires'].get('agl_rank', 'N/A')}")


if __name__ == "__main__":
    sys.exit(main())
